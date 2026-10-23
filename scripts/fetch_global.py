"""Tải số liệu vĩ mô quốc tế từ nguồn PUBLIC (FRED CSV, NY Fed Markets API, DBnomics) → data/global/series.json.

Chạy trên GitHub Actions (máy chủ có internet), không cần API key.
Mỗi khái niệm có danh sách mã ứng viên; script lấy mã có số mới nhất và ghi báo cáo vào data/global/report.json.
Chuỗi lưu từ 2010 để file gọn.
"""
import csv, io, json, sys, time, urllib.parse, urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/global'
START = '2010-01-01'
UA = {'User-Agent': 'Mozilla/5.0 (FinSuccess dashboard; +https://github.com/mdnguyen-cyber/bank-dashboard)'}

# khái niệm: (nhãn, [ứng viên]); ứng viên = 'fred:ID' | 'dbn:provider/dataset/series' | 'nyfed:sofr99'
CAT = {
    # ---------------- Mỹ · lạm phát
    'us_cpi': ('CPI (chỉ số)', ['fred:CPIAUCSL']),
    'us_cpi_core': ('CPI lõi', ['fred:CPILFESL']),
    'us_cpi_shelter': ('CPI nhà ở (shelter)', ['fred:CUSR0000SAH1']),
    'us_cpi_coregoods': ('CPI hàng hóa lõi', ['fred:CUSR0000SACL1E']),
    'us_cpi_svc_ex_energy': ('CPI dịch vụ trừ năng lượng', ['fred:CUSR0000SASLE']),
    'us_cpi_supercore': ('CPI dịch vụ trừ nhà ở (supercore)', ['fred:CUSR0000SASL2RS']),
    'us_cpi_energy': ('CPI năng lượng', ['fred:CPIENGSL']),
    'us_cpi_food': ('CPI lương thực', ['fred:CPIUFDSL']),
    'us_pce': ('PCE', ['fred:PCEPI']),
    'us_pce_core': ('PCE lõi', ['fred:PCEPILFE']),
    'us_pce_supercore': ('PCE dịch vụ trừ năng lượng và nhà ở (supercore)', ['fred:IA001260M']),
    'us_pce_goods': ('PCE hàng hóa', ['fred:DGDSRG3M086SBEA']),
    'us_pce_services': ('PCE dịch vụ', ['fred:DSERRG3M086SBEA']),
    'us_pce_trim': ('PCE trimmed mean 12 tháng (Dallas Fed)', ['fred:PCETRIM12M159SFRBDAL']),
    'us_ppi_fd': ('PPI cầu cuối cùng', ['fred:PPIFIS', 'fred:WPSFD49207']),
    'us_ppi_core': ('PPI cầu cuối cùng trừ thực phẩm, năng lượng, thương mại', ['fred:PPICOR', 'fred:WPSFD4111']),
    'us_ppi_inter': ('PPI trung gian (hàng chế biến)', ['fred:WPSID61']),
    'us_ppi_crude': ('PPI nguyên liệu thô', ['fred:WPSID62']),
    'us_breakeven5': ('Lạm phát kỳ vọng 5 năm (breakeven)', ['fred:T5YIE']),
    # ---------------- Mỹ · lao động, tăng trưởng
    'us_unrate': ('Tỷ lệ thất nghiệp', ['fred:UNRATE']),
    'us_payems': ('Việc làm phi nông nghiệp (nghìn)', ['fred:PAYEMS']),
    'us_ahe': ('Thu nhập bình quân/giờ', ['fred:CES0500000003']),
    'us_claims': ('Đơn xin trợ cấp thất nghiệp lần đầu', ['fred:ICSA']),
    'us_jolts': ('Số việc làm trống (JOLTS)', ['fred:JTSJOL']),
    'us_gdp': ('GDP thực', ['fred:GDPC1']),
    # ---------------- Khảo sát: niềm tin, kỳ vọng lạm phát (FRED, tự cập nhật). PMI ISM/S&P/NBS không có nguồn công khai cập nhật (DBnomics dừng 2025–02/2026) → data/backfill/surveys.json
    'us_umcsent': ('Niềm tin tiêu dùng ĐH Michigan', ['fred:UMCSENT']),
    'us_mich': ('Kỳ vọng lạm phát 1 năm ĐH Michigan (%)', ['fred:MICH']),
    'us_gdppot': ('GDP tiềm năng (CBO)', ['fred:GDPPOT']),
    'us_nrou': ('Thất nghiệp tự nhiên (NAIRU, CBO)', ['fred:NROU']),
    # ---------------- Mỹ · Fed, lãi suất, thanh khoản
    'us_effr': ('Lãi suất quỹ liên bang hiệu dụng', ['fred:DFF']),
    'us_fedtgt_up': ('Cận trên mục tiêu Fed', ['fred:DFEDTARU']),
    'us_fedtgt_lo': ('Cận dưới mục tiêu Fed', ['fred:DFEDTARL']),
    'us_iorb': ('IORB', ['fred:IORB']),
    'us_sofr': ('SOFR', ['fred:SOFR']),
    'us_tgcr': ('TGCR', ['fred:TGCRRATE']),
    'us_sofr99': ('SOFR phân vị 99', ['nyfed:sofr99']),
    'us_2y': ('Lợi suất TPCP Mỹ 2 năm', ['fred:DGS2']),
    'us_10y': ('Lợi suất TPCP Mỹ 10 năm', ['fred:DGS10']),
    'us_walcl': ('Tổng tài sản Fed (triệu USD)', ['fred:WALCL']),
    'us_bills': ('Fed nắm giữ tín phiếu Kho bạc (triệu USD)', ['fred:WSHOBL']),
    'us_rstar': ('Lãi suất trung lập r* (HLW, NY Fed)', ['hlw:US']),
    'us_reserves': ('Dự trữ của NH tại Fed (triệu USD)', ['fred:WRESBAL']),
    'us_tga': ('Tài khoản Kho bạc tại Fed - TGA (triệu USD)', ['fred:WTREGEN']),
    'us_rrp': ('ON RRP (tỷ USD)', ['fred:RRPONTSYD']),
    'us_repo': ('Repo thường trực SRF (tỷ USD)', ['fred:RPONTSYD']),
    'us_dxy': ('Chỉ số USD danh nghĩa rộng', ['fred:DTWEXBGS']),
    'us_casehiller': ('Giá nhà Case-Shiller', ['fred:CSUSHPISA']),
    'us_wti': ('Dầu WTI', ['fred:DCOILWTICO']),
    'us_brent': ('Dầu Brent', ['fred:DCOILBRENTEU']),
    # ---------------- Mỹ · cấu phần bảng cân đối Fed (H.4.1, triệu USD, theo tuần)
    'fed_notes': ('Fed nắm giữ trái phiếu Kho bạc kỳ hạn dài (notes & bonds danh nghĩa)', ['fred:WSHONBNL']),
    'fed_tips': ('Fed nắm giữ TIPS', ['fred:WSHONBIIL']),
    'fed_tips_ic': ('Bù lạm phát của TIPS Fed nắm giữ', ['fred:WSHOICL']),
    'fed_ust': ('Fed nắm giữ tổng TPCP Mỹ', ['fred:WSHOTSL']),
    'fed_mbs': ('Fed nắm giữ MBS', ['fred:WSHOMCB']),
    'fed_loans': ('Cho vay của Fed (tổng)', ['fred:WLCFLL']),
    'fed_currency': ('Tiền mặt lưu hành', ['fred:WCURCIR']),
    'fed_rrp_all': ('Reverse repo (tổng, gồm nước ngoài)', ['fred:WLRRAL']),
    'fed_res_wed': ('Dự trữ NH tại Fed, số ngày thứ Tư', ['fred:WRBWFRBL']),
    'fed_tga_wed': ('TGA, số ngày thứ Tư', ['fred:WDTGAL']),
    'fed_currency_wed': ('Tiền mặt lưu hành, số ngày thứ Tư', ['fred:WLFN', 'fred:WCURCIR']),
    'ust_nfc': ('UST do doanh nghiệp phi tài chính nắm', ['fred:BOGZ1LM103061103Q', 'fred:BOGZ1FL103061103Q']),
    'ust_etf': ('UST do ETF nắm', ['fred:BOGZ1LM563061103Q', 'fred:BOGZ1FL563061103Q']),
    'ust_gse': ('UST do GSE nắm', ['fred:BOGZ1FL403061105Q', 'fred:BOGZ1LM403061105Q']),
    'ust_slpen': ('UST do quỹ hưu trí công chức bang, địa phương nắm', ['fred:BOGZ1FL223061143Q', 'fred:BOGZ1LM223061143Q']),
    'ust_fedpen': ('UST do quỹ hưu trí công chức liên bang nắm', ['fred:BOGZ1FL343061105Q', 'fred:BOGZ1LM343061105Q']),
    'fed_rrp_foreign': ('Reverse repo của tổ chức nước ngoài', ['fred:WLRRAFOIAL']),
    # ---------------- Mỹ · Kho bạc: nợ, thâm hụt, ai nắm UST (Z.1 theo quý, triệu USD; TIC tỷ USD)
    'tsy_debt_total': ('Tổng nợ công liên bang (triệu USD)', ['fred:GFDEBTN']),
    'tsy_debt_public': ('Nợ liên bang do công chúng nắm giữ (triệu USD)', ['fred:FYGFDPUN']),
    'tsy_deficit': ('Thặng dư (+)/thâm hụt (−) ngân sách tháng (triệu USD)', ['fred:MTSDS133FMS']),
    'ust_total': ('TPCP Mỹ lưu hành (Z.1)', ['fred:BOGZ1FL313161105Q', 'fred:BOGZ1FL313161100Q']),
    'ust_fed': ('UST do Fed nắm', ['fred:BOGZ1FL713061103Q', 'fred:BOGZ1FL713061100Q', 'fred:FDHBFRBN']),
    'ust_row': ('UST do nước ngoài nắm', ['fred:BOGZ1FL263061105Q', 'fred:BOGZ1LM263061105Q', 'fred:BOGZ1FL263061103Q', 'fred:FDHBFIN']),
    'ust_row_official': ('UST do tổ chức chính thức nước ngoài nắm', ['fred:BOGZ1FL263061120Q', 'fred:FDHBFIN']),
    'ust_banks': ('UST do ngân hàng Mỹ nắm', ['fred:BOGZ1FL763061100Q', 'fred:BOGZ1FL763061105Q', 'fred:TLTSYSUS']),
    'ust_mmf': ('UST do quỹ thị trường tiền tệ nắm', ['fred:BOGZ1FL633061105Q', 'fred:BOGZ1FL633061100Q']),
    'ust_mf': ('UST do quỹ tương hỗ nắm', ['fred:BOGZ1FL653061105Q', 'fred:BOGZ1FL653061100Q']),
    'ust_hh': ('UST do hộ gia đình & tổ chức phi lợi nhuận nắm (gồm quỹ phòng hộ)', ['fred:BOGZ1FL153061105Q', 'fred:BOGZ1LM153061105Q', 'fred:BOGZ1FL153061103Q', 'fred:BOGZ1LM153061103Q']),
    'ust_ins': ('UST do bảo hiểm nắm', ['fred:BOGZ1FL513061105Q', 'fred:BOGZ1FL543061105Q']),
    'ust_pension': ('UST do quỹ hưu trí nắm', ['fred:BOGZ1FL593061105Q', 'fred:BOGZ1FL573061105Q']),
    'ust_slg': ('UST do chính quyền bang & địa phương nắm', ['fred:BOGZ1FL213061103Q', 'fred:BOGZ1LM213061103Q', 'fred:BOGZ1FL213061105Q']),
    'ust_broker': ('UST do công ty môi giới chứng khoán nắm', ['fred:BOGZ1FL663061105Q', 'fred:BOGZ1FL663061100Q']),
    # ---------------- Khu vực đồng Euro
    'ea_hicp': ('HICP (% YoY)', ['ecb:HICP/M.U2.N.000000.4D0.ANR', 'dbn:ECB/HICP/M.U2.N.000000.4D0.ANR', 'ecb:ICP/M.U2.N.000000.4.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.EA20', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.EA']),
    'ea_hicp_core': ('HICP lõi trừ năng lượng, thực phẩm (% YoY)', ['ecb:HICP/M.U2.N.XEF000.4D0.ANR', 'dbn:ECB/HICP/M.U2.N.XEF000.4D0.ANR', 'ecb:ICP/M.U2.N.XEF000.4.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.TOT_X_NRG_FOOD.EA20', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.TOT_X_NRG_FOOD.EA']),
    'ea_hicp_serv': ('HICP dịch vụ (% YoY)', ['ecb:HICP/M.U2.N.SERV00.4D0.ANR', 'dbn:ECB/HICP/M.U2.N.SERV00.4D0.ANR', 'ecb:ICP/M.U2.N.SERV00.4.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.SERV.EA20', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.SERV.EA']),
    'ea_dfr': ('Lãi suất tiền gửi ECB', ['fred:ECBDFR']),
    'ea_assets': ('Tổng tài sản Eurosystem (triệu EUR)', ['fred:ECBASSETSW']),
    'ea_unrate': ('Thất nghiệp khu vực Euro', ['ecb:LFSI/M.I9.S.UNEHRT.TOTAL0.15_74.T', 'fred:LRHUTTTTEZM156S']),
    'de_10y': ('Lợi suất Bund 10 năm', ['fred:IRLTLT01DEM156N']),
    'eurusd': ('EUR/USD', ['fred:DEXUSEU']),
    # ---------------- Nhật Bản
    'jp_cpi_idx': ('CPI Nhật (chỉ số 2020=100)', ['dbn:STATJP/CPIm/001']),
    'jp_cpi_xff_idx': ('CPI Nhật trừ thực phẩm tươi sống (chỉ số)', ['dbn:STATJP/CPIm/733']),
    'jp_cpi_xffe_idx': ('CPI Nhật trừ thực phẩm tươi sống và năng lượng (chỉ số)', ['dbn:STATJP/CPIm/740']),
    'jp_cpi': ('CPI Nhật (% YoY)', ['oecd:JPN.M.N.CPI.PA._T.N.GY', 'dbn:IMF/CPI/M.JP.PCPI_PC_CP_A_PT', 'fred:CPALTT01JPM659N']),
    'jp_cpi_core': ('CPI Nhật trừ thực phẩm và năng lượng (% YoY)', ['oecd:JPN.M.N.CPI.PA._TXCP01_NRG.N.GY', 'fred:CPGRLE01JPM659N', 'dbn:OECD/DSD_PRICES@DF_PRICES_ALL/JPN.M.N.CPI.PA._TXCP01_NRG.N.GY']),
    'jp_rate': ('Lãi suất qua đêm Nhật', ['fred:IRSTCI01JPM156N', 'fred:IR3TIB01JPM156N']),
    'jp_assets': ('Tổng tài sản BOJ (100 triệu JPY)', ['fred:JPNASSETS']),
    'jp_10y': ('Lợi suất JGB 10 năm', ['fred:IRLTLT01JPM156N']),
    'usdjpy': ('USD/JPY', ['fred:DEXJPUS']),
    'jp_reserves': ('Dự trữ ngoại hối Nhật trừ vàng (triệu USD)', ['fred:TRESEGJPM052N']),
    # ---------------- Toàn cầu: lợi suất 10 năm theo tháng (ECB — lãi suất dài hạn hội tụ; OECD qua FRED)
    'y10_de': ('Đức 10 năm', ['ecb:IRS/M.DE.L.L40.CI.0000.EUR.N.Z', 'fred:IRLTLT01DEM156N']),
    'y10_fr': ('Pháp 10 năm', ['ecb:IRS/M.FR.L.L40.CI.0000.EUR.N.Z', 'fred:IRLTLT01FRM156N']),
    'y10_it': ('Ý 10 năm', ['ecb:IRS/M.IT.L.L40.CI.0000.EUR.N.Z', 'fred:IRLTLT01ITM156N']),
    'y10_es': ('Tây Ban Nha 10 năm', ['ecb:IRS/M.ES.L.L40.CI.0000.EUR.N.Z', 'fred:IRLTLT01ESM156N']),
    'y10_nl': ('Hà Lan 10 năm', ['ecb:IRS/M.NL.L.L40.CI.0000.EUR.N.Z', 'fred:IRLTLT01NLM156N']),
    'y10_gb': ('Anh 10 năm', ['fred:IRLTLT01GBM156N']),
    'y10_kr': ('Hàn Quốc 10 năm', ['fred:IRLTLT01KRM156N']),
    'y10_au': ('Úc 10 năm', ['fred:IRLTLT01AUM156N']),
    'y10_ca': ('Canada 10 năm', ['fred:IRLTLT01CAM156N']),
    # ---------------- Toàn cầu: CPI % YoY theo tháng
    'cpi_de': ('CPI Đức', ['ecb:HICP/M.DE.N.000000.4D0.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.DE', 'oecd:DEU.M.N.CPI.PA._T.N.GY']),
    'cpi_it': ('CPI Ý', ['ecb:HICP/M.IT.N.000000.4D0.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.IT', 'oecd:ITA.M.N.CPI.PA._T.N.GY']),
    'cpi_es': ('CPI Tây Ban Nha', ['ecb:HICP/M.ES.N.000000.4D0.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.ES', 'oecd:ESP.M.N.CPI.PA._T.N.GY']),
    'cpi_fr': ('CPI Pháp', ['ecb:HICP/M.FR.N.000000.4D0.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.FR', 'oecd:FRA.M.N.CPI.PA._T.N.GY']),
    'cpi_nl': ('CPI Hà Lan', ['ecb:HICP/M.NL.N.000000.4D0.ANR', 'dbn:Eurostat/prc_hicp_manr/M.RCH_A.CP00.NL', 'oecd:NLD.M.N.CPI.PA._T.N.GY']),
    'cpi_kr': ('CPI Hàn Quốc', ['bis:KR', 'oecd:KOR.M.N.CPI.PA._T.N.GY', 'dbn:IMF/CPI/M.KR.PCPI_PC_CP_A_PT']),
    'cpi_in': ('CPI Ấn Độ', ['bis:IN', 'oecd:IND.M.N.CPI.PA._T.N.GY', 'dbn:IMF/CPI/M.IN.PCPI_PC_CP_A_PT']),
    'cpi_id': ('CPI Indonesia', ['bis:ID', 'oecd:IDN.M.N.CPI.PA._T.N.GY', 'dbn:IMF/CPI/M.ID.PCPI_PC_CP_A_PT']),
    'cpi_sg': ('CPI Singapore', ['bis:SG', 'dbn:IMF/CPI/M.SG.PCPI_PC_CP_A_PT', 'oecd:SGP.M.N.CPI.PA._T.N.GY']),
    'cpi_th': ('CPI Thái Lan', ['bis:TH', 'dbn:IMF/CPI/M.TH.PCPI_PC_CP_A_PT', 'oecd:THA.M.N.CPI.PA._T.N.GY']),
    'cpi_my': ('CPI Malaysia', ['bis:MY', 'dbn:IMF/CPI/M.MY.PCPI_PC_CP_A_PT', 'oecd:MYS.M.N.CPI.PA._T.N.GY']),
    'cpi_ph': ('CPI Philippines', ['bis:PH', 'dbn:IMF/CPI/M.PH.PCPI_PC_CP_A_PT', 'oecd:PHL.M.N.CPI.PA._T.N.GY']),
    'cpi_tw': ('CPI Đài Loan', ['bis:TW', 'oecd:TWN.M.N.CPI.PA._T.N.GY']),
    # ---------------- Trung Quốc
    'cn_cpi': ('CPI Trung Quốc (% YoY)', ['oecd:CHN.M.N.CPI.PA._T.N.GY', 'dbn:IMF/CPI/M.CN.PCPI_PC_CP_A_PT', 'fred:CPALTT01CNM659N']),
    'cn_ppi': ('PPI Trung Quốc (% YoY)', ['nbs:NBS/M_A010B/A010B01|NBS/M_A010H/A010H01']),
    'cn_rate': ('Lãi suất liên ngân hàng Trung Quốc', ['fred:IR3TIB01CNM156N', 'fred:INTDSRCNM193N']),
    'usdcny': ('USD/CNY', ['fred:DEXCHUS']),
}


def get(url, tries=3):
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
                return r.read().decode('utf-8')
        except Exception as e:  # noqa: BLE001
            err = e
            time.sleep(2 * (i + 1))
    raise err


def fred(sid):
    txt = get(f'https://fred.stlouisfed.org/graph/fredgraph.csv?id={sid}&cosd={START}')
    rows = list(csv.reader(io.StringIO(txt)))
    out = []
    for d, v in rows[1:]:
        if v not in ('', '.'):
            out.append([d, round(float(v), 4)])
    return out


def dbn(path):
    js = json.loads(get(f'https://api.db.nomics.world/v22/series/{path}?observations=1&format=json'))
    doc = js['series']['docs'][0]
    out = []
    for p, v in zip(doc['period'], doc['value']):
        if v not in (None, 'NA') and p >= START[:7]:
            d = p + '-01' if len(p) == 7 else p
            out.append([d, round(float(v), 4)])
    return out


def ecb(key):
    flow, k = key.split('/', 1)
    txt = get(f'https://data-api.ecb.europa.eu/service/data/{flow}/{k}?format=csvdata&startPeriod={START[:7]}')
    rows = list(csv.DictReader(io.StringIO(txt)))
    return [[(r['TIME_PERIOD'] + '-01') if len(r['TIME_PERIOD']) == 7 else r['TIME_PERIOD'], round(float(r['OBS_VALUE']), 4)] for r in rows if r.get('OBS_VALUE') not in (None, '', 'NaN')]


def oecd(key):
    url = f'https://sdmx.oecd.org/public/rest/data/OECD.SDD.TPS,DSD_PRICES@DF_PRICES_ALL,1.0/{key}?startPeriod={START[:7]}&format=csvfile'
    rows = list(csv.DictReader(io.StringIO(get(url))))
    out = [[r['TIME_PERIOD'] + '-01', round(float(r['OBS_VALUE']), 4)] for r in rows if r.get('OBS_VALUE') not in (None, '')]
    return sorted(out)


PROBES = ['ISM manufacturing PMI', 'ISM services PMI', 'manufacturing purchasing managers index China', 'non-manufacturing PMI China', 'economic sentiment indicator euro area', 'purchasing managers index']


def probe():
    out = {}
    for q in PROBES:
        try:
            js = json.loads(get('https://api.db.nomics.world/v22/search?q=' + urllib.parse.quote(q) + '&limit=8'))
            out[q] = [{'code': f"{d.get('provider_code')}/{d.get('code')}", 'name': d.get('name'), 'nb': d.get('nb_matching_series')} for d in js['results']['docs']]
        except Exception as e:  # noqa: BLE001
            out[q] = str(e)[:200]
    try:
        js = json.loads(get('https://api.db.nomics.world/v22/series/ECB/HICP?q=' + urllib.parse.quote('monthly euro area changing composition annual rate of change') + '&limit=40&observations=0'))
        out['ECB/HICP M'] = [(d['series_code'], d.get('series_name', '')[:120]) for d in js['series']['docs'] if d['series_code'].startswith('M.U2')]
    except Exception as e:  # noqa: BLE001
        out['ECB/HICP M'] = str(e)[:200]
    return out


def nbs(paths):   # NBS 'cùng tháng năm trước = 100' → % YoY, nối các giai đoạn
    out = {}
    for pth in paths.split('|'):
        for d, v in dbn(pth):
            out[d] = round(v - 100, 2)
    return sorted(out.items())


def hlw(country):   # Holston-Laubach-Williams r* (NY Fed), file xlsx công khai
    import openpyxl
    url = 'https://www.newyorkfed.org/medialibrary/media/research/economists/williams/data/Holston_Laubach_Williams_current_estimates.xlsx'
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
        wb = openpyxl.load_workbook(io.BytesIO(r.read()), read_only=True, data_only=True)
    ws = wb['HLW Estimates'] if 'HLW Estimates' in wb.sheetnames else wb.worksheets[0]
    out = []
    for row in ws.iter_rows(values_only=True):
        d, v = row[0], row[10] if len(row) > 10 else None   # cột K = r* Mỹ (Date, …, Natural Rate US)
        if hasattr(d, 'year') and isinstance(v, (int, float)) and d.year >= 2010:
            out.append([d.strftime('%Y-%m-%d'), round(float(v), 4)])
    return out


def nyfed_sofr99():
    js = json.loads(get(f'https://markets.newyorkfed.org/api/rates/secured/sofr/search.json?startDate=2018-04-02&endDate=2100-01-01'))
    out = []
    for r in js['refRates']:
        try:
            out.append([r['effectiveDate'], float(r.get('percentPercentile99'))])
        except (TypeError, ValueError):
            pass
    return sorted(out)


def bis(cc):   # BIS: CPI % YoY (WS_LONG_CPI, 771 = thay đổi so cùng kỳ)
    txt = get(f'https://stats.bis.org/api/v1/data/WS_LONG_CPI/M.{cc}.771?format=csv&startPeriod={START[:7]}')
    rows = list(csv.DictReader(io.StringIO(txt)))
    return sorted([[r['TIME_PERIOD'] + '-01', round(float(r['OBS_VALUE']), 4)] for r in rows if r.get('OBS_VALUE') not in (None, '', 'NaN')])


def fetch(c):
    src, sid = c.split(':', 1)
    return {'fred': fred, 'dbn': dbn, 'ecb': ecb, 'oecd': oecd, 'bis': bis, 'nbs': nbs, 'hlw': hlw, 'nyfed': lambda _: nyfed_sofr99()}[src](sid)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    old = {}
    p = OUT / 'series.json'
    if p.exists():
        old = json.loads(p.read_text(encoding='utf-8')).get('series', {})
    series, report = {}, {}
    for key, (label, cands) in CAT.items():
        best, tried = None, []
        for c in cands:
            try:
                obs = fetch(c)
                tried.append({'id': c, 'n': len(obs), 'last': obs[-1][0] if obs else None})
                if obs and (best is None or obs[-1][0] > best[1][-1][0]):
                    best = (c, obs)
            except Exception as e:  # noqa: BLE001
                tried.append({'id': c, 'error': str(e)[:160]})
        if best:
            series[key] = {'label': label, 'src': best[0], 'obs': best[1]}
        elif key in old:   # giữ bản cũ nếu lần này lỗi mạng
            series[key] = old[key]
            series[key]['stale'] = True
        report[key] = tried
        print(key, '→', best[0] if best else 'KHÔNG CÓ', best[1][-1] if best else '', file=sys.stderr)
    stamp = time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
    (OUT / 'series.json').write_text(json.dumps({'fetched': stamp, 'series': series}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    (OUT / 'report.json').write_text(json.dumps({'fetched': stamp, 'report': report}, ensure_ascii=False, indent=1), encoding='utf-8')
    try:
        (OUT / 'probe.json').write_text(json.dumps(probe(), ensure_ascii=False, indent=1), encoding='utf-8')
    except Exception as e:  # noqa: BLE001
        print('probe lỗi', e, file=sys.stderr)
    miss = [k for k in CAT if k not in series]
    print(f'{len(series)}/{len(CAT)} chuỗi; thiếu: {miss}', file=sys.stderr)


if __name__ == '__main__':
    main()
