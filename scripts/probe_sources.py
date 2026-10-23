"""Dò nguồn công khai (chạy trên GitHub Actions) → data/probe/*.json: mẫu phản hồi, danh sách link, để chọn endpoint."""
import urllib.parse, json, os, re, sys, urllib.request
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / 'data/probe'
UA = {'User-Agent': 'Mozilla/5.0 (FinSuccess dashboard; +https://github.com/mdnguyen-cyber/bank-dashboard)'}
FD = 'https://api.fiscaldata.treasury.gov/services/api/fiscal_service'
BOJ = 'https://www.stat-search.boj.or.jp/api/v1'
URLS = {
    'tcbs_bars': 'https://apipubaws.tcbs.com.vn/stock-insight/v2/stock/bars-long-term?ticker=VCB&type=stock&resolution=D&countBack=30&to=1792800000',
    'tcbs_ratio': 'https://apipubaws.tcbs.com.vn/tcanalysis/v1/finance/VCB/financialratio?yearly=0&isAll=false',
    'tcbs_overview': 'https://apipubaws.tcbs.com.vn/tcanalysis/v1/ticker/VCB/overview',
    'vnd_prices': 'https://api-finfo.vndirect.com.vn/v4/stock_prices?sort=date&q=code:VCB&size=5',
    'vnd_ratios': 'https://api-finfo.vndirect.com.vn/v4/ratios/latest?filter=ratioCode:PRICE_TO_BOOK,PRICE_TO_EARNINGS&where=code:VCB&order=reportDate&fields=ratioCode,value,reportDate',
    'vnd_fin': 'https://api-finfo.vndirect.com.vn/v4/financial_statements?q=code:VCB~reportType:QUARTER~modelType:1&sort=fiscalDate&size=5',
    'ssi_chart': 'https://iboard-api.ssi.com.vn/statistics/charts/history?resolution=1D&symbol=VCB&from=1790000000&to=1792800000',
    'vps_chart': 'https://histdatafeed.vps.com.vn/tradingview/history?symbol=VCB&resolution=D&from=1790000000&to=1792800000',
    'yahoo_gold': 'https://query1.finance.yahoo.com/v8/finance/chart/GC=F?range=1mo&interval=1d',
    'cafef_rss_nh': 'https://cafef.vn/tai-chinh-ngan-hang.rss',
    'vnexpress_rss_kd': 'https://vnexpress.net/rss/kinh-doanh.rss',
    'sbv_home_raw': 'https://www.sbv.gov.vn/',
    'hnx_yield_raw': 'https://hnx.vn/vi-vn/trai-phieu/duong-cong-loi-suat.html',
    'vbma_raw': 'https://vbma.org.vn/vi/market-data/yield-curve',
}
KEEP = re.compile(r'government|bond|yield', re.I)
KEEP = re.compile(r'government|bond|yield', re.I)
KEEP = re.compile(r'government|JGB|Treasury|reserve|Bank of Japan Accounts|current account|monetary base|intervention|ETF|loans', re.I)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rep = {}
    for k, u in URLS.items():
        try:
            data = None
            if isinstance(u, tuple):
                u, data = u[0], urllib.parse.urlencode(u[1]).encode()
            with urllib.request.urlopen(urllib.request.Request(u, data=data, headers={**UA, 'X-Requested-With': 'XMLHttpRequest', 'Referer': 'https://www.tpex.org.tw/en-us/bond/info/statistics-gb/day/ebts.html'}), timeout=60) as r:
                b = r.read(); ct = r.headers.get('Content-Type', '')
            if u.endswith('.xlsx'):
                import openpyxl, io
                wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
                for ws in wb.worksheets: ws.reset_dimensions()
                rep[k] = {'url': u, 'status': 200, 'bytes': len(b), 'sheets': {ws.title: [[str(c) for c in r] for i, r in enumerate(ws.iter_rows(values_only=True)) if i < 25] for ws in wb.worksheets}}
                rows = list(wb.worksheets[0].iter_rows(values_only=True)); rep[k]['tail'] = [[str(c) for c in r] for r in rows[-15:]]
                continue
            if u.endswith('.xls'):
                import xlrd
                wb = xlrd.open_workbook(file_contents=b)
                rep[k] = {'url': u, 'status': 200, 'sheets': {sh.name: [[str(c.value) for c in sh.row(i)] for i in range(min(sh.nrows, 60))] for sh in wb.sheets()}}
                continue
            if k.endswith('_pdf'):
                (OUT / (k + '.pdf')).write_bytes(b); rep[k] = {'url': u, 'status': 200, 'bytes': len(b)}; continue
            txt = b.decode('utf-8', 'replace') if 'pdf' not in ct else ''
            info = {'url': u, 'status': 200, 'ct': ct, 'bytes': len(b)}
            if k.endswith('_raw'):
                info['tail'] = txt[txt.find('modifyJson') - 200:][:5000] if 'modifyJson' in txt else txt[-5000:]
                info['form'] = re.findall(r'<form[^>]*>', txt)[:5]
                info['urls'] = sorted(set(re.findall(r'https?://[^"\'\s<>)]+', txt)))[:150]
                info['paths'] = sorted(set(re.findall(r'["\'](/(?:api|_next/data|embed|nhung)[^"\'\s]*)', txt)))[:80]
                info['scripts'] = [m[:1500] for m in re.findall(r'<script[^>]*>(.*?)</script>', txt, re.S) if len(m) > 200][:8]
                info['size'] = len(txt)
                info['ctx2'] = [txt[max(0, m.start() - 600):m.end() + 600] for m in re.finditer(r'xlsx|csv|download|probability', txt, re.I)][:6]
            if k.startswith('tw'):
                info['ctx'] = [txt[max(0, m.start() - 300):m.end() + 300] for m in re.finditer(r'BDdys100|Benchmark|storage/|/www/|ajax|fetch\(|\$\.(?:get|post)', txt)][:25]
            if k == 'tw_tpex_api_list':
                js = json.loads(txt); info['paths'] = {pth: (v.get('get') or {}).get('summary') for pth, v in js.get('paths', {}).items()}
            elif k.startswith('boj_meta'):
                js = json.loads(txt)
                info['series'] = [[x['SERIES_CODE'], x['NAME_OF_TIME_SERIES'][:160], x['FREQUENCY'], x['END_OF_THE_TIME_SERIES'], x['UNIT']] for x in js.get('RESULTSET', [])
                                  if x['SERIES_CODE'] and '(Discontinued)' not in x['NAME_OF_TIME_SERIES'] and (k != 'boj_meta_FF' or KEEP.search(x['NAME_OF_TIME_SERIES']))][:1500]
                info['n'] = len(js.get('RESULTSET', [])); info['msg'] = js.get('MESSAGE')
            elif txt.lstrip().startswith('{'):
                js = json.loads(txt)
                info['json'] = {kk: (vv[:20] if isinstance(vv, list) else vv) for kk, vv in js.items()} if isinstance(js, dict) else js
                if isinstance(js, dict) and 'meta' in js:
                    info['json']['meta'] = js['meta']
                s = json.dumps(info['json'], ensure_ascii=False)
                if len(s) > 60000:
                    info['json'] = s[:60000]
            elif '<html' in txt[:2000].lower() or '<!doctype' in txt[:200].lower():
                info['links'] = sorted(set(re.findall(r'href="([^"]+\.(?:csv|xlsx?|pdf|htm|html|txt|zip))"', txt)))[:400]
                info['xls'] = sorted(set(re.findall(r'[\w/.:-]*(?:BDdys|Curve|govbond|benchmark|/www/|\.xls|\.js|\.csv)[\w/.?=&-]*', txt, re.I)))[:200]
                info['api'] = sorted(set(re.findall(r'''["'](/(?:www|api|openapi)/[^"'\s]+)["']''', txt)))[:100]
                info['text'] = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', re.sub(r'(?s)<(script|style).*?</\1>', ' ', txt)))[:3000]
                info['title'] = (re.search(r'<title>(.*?)</title>', txt, re.S) or [None, ''])[1].strip()[:200]
            else:
                info['head'] = txt[:6000]
            rep[k] = info
        except Exception as e:  # noqa: BLE001
            rep[k] = {'url': u, 'err': str(e)[:300]}
        print(k, rep[k].get('status', rep[k].get('err')), file=sys.stderr)
    (OUT / os.environ.get('PROBE_OUT', 'probe2.json')).write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding='utf-8')


if __name__ == '__main__':
    main()
